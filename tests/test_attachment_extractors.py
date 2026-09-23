from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter

from app.extractors.attachments import DocxExtractor, ImageOcrExtractor, PdfExtractor, XlsxExtractor
from app.schemas.attachments import DocumentType


def make_pdf() -> bytes:
    buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(buffer)
    return buffer.getvalue()


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Cable specification")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "article"
    table.cell(0, 1).text = "A-1"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Products"
    worksheet.append(["article", "quantity"])
    worksheet.append(["A-1", 3])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_image(image_format: str = "PNG") -> bytes:
    buffer = BytesIO()
    Image.new("RGB", (20, 20), "white").save(buffer, format=image_format)
    return buffer.getvalue()


def test_pdf_extractor_marks_blank_text_pdf_for_ocr() -> None:
    result = PdfExtractor().extract("scan.pdf", "application/pdf", make_pdf())

    assert result.document_type == DocumentType.PDF
    assert result.metadata["page_count"] == 1
    assert "empty_document" in result.warnings


def test_docx_extractor_returns_text_and_tables() -> None:
    result = DocxExtractor().extract("spec.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", make_docx())

    assert result.document_type == DocumentType.DOCX
    assert result.text == "Cable specification"
    assert result.tables[0].rows[0] == ["article", "A-1"]


def test_xlsx_extractor_returns_structured_sheet_rows() -> None:
    result = XlsxExtractor(max_sheets=2, max_rows=10, max_columns=5, max_uncompressed_bytes=1_000_000).extract(
        "items.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", make_xlsx()
    )

    assert result.document_type == DocumentType.XLSX
    assert result.tables[0].name == "Products"
    assert result.tables[0].rows[1] == ["A-1", 3]


def test_xlsx_extractor_limits_rows() -> None:
    result = XlsxExtractor(max_sheets=2, max_rows=1, max_columns=5, max_uncompressed_bytes=1_000_000).extract(
        "items.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", make_xlsx()
    )

    assert result.tables[0].rows == [["article", "quantity"]]
    assert "xlsx_sheet_truncated:Products" in result.warnings


def test_image_ocr_extractor_uses_tesseract(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("app.extractors.attachments.pytesseract.image_to_string", lambda image: "A-1")

    result = ImageOcrExtractor().extract("scan.png", "image/png", make_image())

    assert result.document_type == DocumentType.IMAGE
    assert result.text == "A-1"


def test_image_ocr_failure_returns_warning(monkeypatch: pytest.MonkeyPatch) -> None:
    import pytesseract

    def fail(image):
        raise pytesseract.TesseractNotFoundError()

    monkeypatch.setattr("app.extractors.attachments.pytesseract.image_to_string", fail)
    result = ImageOcrExtractor().extract("scan.png", "image/png", make_image())

    assert result.text == ""
    assert "ocr_failed" in result.warnings
