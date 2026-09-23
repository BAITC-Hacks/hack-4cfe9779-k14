from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from uuid import UUID, uuid4

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.schemas.catalog import CatalogCandidate, CatalogSearchResponse, CurrentAvailability
from app.schemas.chat import ChatAnalysis, ChatIntent, ChatMessageCreate, ChatRole
from app.services.attachments import AttachmentService
from app.services.chat import ChatService

pytestmark = pytest.mark.asyncio


class FakeChatRepository:
    def __init__(self) -> None:
        self.session_id = uuid4()
        self.messages = []

    async def get_session(self, session_id: UUID):
        return SimpleNamespace(id=session_id) if session_id == self.session_id else None

    async def add_message(self, session_id, role, content, metadata=None):
        message = SimpleNamespace(
            id=uuid4(), session_id=session_id, role=role.value, content=content,
            created_at=datetime.now(timezone.utc), metadata_json=metadata or {},
        )
        self.messages.append(message)
        return message

    async def history(self, session_id, *, limit=None):
        messages = [message for message in self.messages if message.session_id == session_id]
        return messages[-limit:] if limit else messages


class FakeAttachmentRepository:
    def __init__(self, records) -> None:
        self.records = records

    async def get_for_session(self, session_id, attachment_ids):
        return [record for record in self.records if record.id in attachment_ids and record.session_id == session_id]


class FakeCatalogService:
    def __init__(self) -> None:
        self.current_articles: list[str] = []
        self.candidate = CatalogCandidate(
            id=uuid4(), article="A-1", external_id="external-A-1", name="Cable", description=None,
            brand="EKT", characteristics={}, cached_price=Decimal("1.00"), cached_stock_by_location={"local": 1},
            cached_available=True,
        )

    async def search_candidates(self, query, *, characteristics=None, limit=20):
        del characteristics, limit
        return CatalogSearchResponse(candidates=[self.candidate] if query == "A-1" else [], match_type="article")

    async def get_current_availability(self, article):
        self.current_articles.append(article)
        return CurrentAvailability(
            article=article, price=Decimal("12.50"), stock_by_location={"warehouse": 10}, available=True, current=True,
        )


class MockLLM:
    def __init__(self) -> None:
        self.attachment_data = None

    async def analyze(self, history, attachment_data=None):
        del history
        self.attachment_data = attachment_data
        return ChatAnalysis(intent=ChatIntent.FIND_PRODUCT, article="A-1")

    async def aclose(self):
        return None


def pdf_bytes() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    contents = DecodedStreamObject()
    contents.set_data(b"BT /F1 12 Tf 40 250 Td (Article: A-1 Quantity: 2) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(contents)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Article: A-1 Quantity: 2")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["article", "quantity"])
    sheet.append(["A-1", 2])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def image_bytes() -> bytes:
    output = BytesIO()
    Image.new("RGB", (20, 20), "white").save(output, format="PNG")
    return output.getvalue()


@pytest.mark.parametrize(
    ("filename", "mime_type", "content"),
    [
        ("items.pdf", "application/pdf", pdf_bytes()),
        ("items.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx_bytes()),
        ("items.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", xlsx_bytes()),
        ("items.png", "image/png", image_bytes()),
    ],
)
async def test_extracted_attachment_is_used_for_catalog_and_current_ekt_check(
    monkeypatch: pytest.MonkeyPatch, filename: str, mime_type: str, content: bytes,
) -> None:
    monkeypatch.setattr("app.extractors.attachments.pytesseract.image_to_string", lambda image: "A-1 2")
    extracted = AttachmentService().process(filename, mime_type, content)
    repository = FakeChatRepository()
    attachment_id = uuid4()
    record = SimpleNamespace(
        id=attachment_id, session_id=repository.session_id, filename=extracted.filename, mime_type=extracted.mime_type,
        document_type=extracted.document_type.value, extracted_text=extracted.text,
        tables=[table.model_dump(mode="json") for table in extracted.tables], warnings=extracted.warnings,
        metadata_json=extracted.metadata,
    )
    catalog = FakeCatalogService()
    llm = MockLLM()
    service = ChatService(repository, catalog, llm, attachments=FakeAttachmentRepository([record]))

    reply = await service.send_message(
        repository.session_id,
        ChatMessageCreate(content="Подбери товары из файла", attachment_ids=[attachment_id]),
    )

    assert llm.attachment_data[0].text == extracted.text
    assert reply.attachment_items[0].article == "A-1"
    assert reply.attachment_items[0].quantity == 2
    assert reply.attachment_items[0].current_data["current"] is True
    assert catalog.current_articles == ["A-1"]
    assert "распознано позиций: 1" in reply.assistant_message.content


async def test_chat_explains_ocr_failure_without_claiming_a_result() -> None:
    repository = FakeChatRepository()
    attachment_id = uuid4()
    record = SimpleNamespace(
        id=attachment_id, session_id=repository.session_id, filename="scan.png", mime_type="image/png",
        document_type="image", extracted_text="", tables=[], warnings=["ocr_failed", "empty_document"], metadata_json={},
    )
    service = ChatService(repository, FakeCatalogService(), MockLLM(), attachments=FakeAttachmentRepository([record]))

    reply = await service.send_message(
        repository.session_id,
        ChatMessageCreate(content="Проверьте скан", attachment_ids=[attachment_id]),
    )

    assert "OCR не смог" in reply.assistant_message.content
    assert reply.attachment_items == []
