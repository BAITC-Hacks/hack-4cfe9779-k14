from __future__ import annotations

from io import BytesIO
from typing import Protocol
from zipfile import BadZipFile, ZipFile

import pytesseract
from docx import Document
from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader

from app.schemas.attachments import AttachmentResult, DocumentType, ExtractedTable


class AttachmentExtractionError(Exception):
    """The uploaded bytes cannot be safely processed as the declared type."""


class AttachmentExtractor(Protocol):
    def extract(self, filename: str, mime_type: str, content: bytes) -> AttachmentResult: ...


class PdfExtractor:
    MIN_TEXT_CHARACTERS = 20
    MAX_PAGES = 200

    def extract(self, filename: str, mime_type: str, content: bytes) -> AttachmentResult:
        try:
            reader = PdfReader(BytesIO(content))
        except Exception as exc:
            raise AttachmentExtractionError("PDF is damaged or unreadable") from exc
        if len(reader.pages) > self.MAX_PAGES:
            raise AttachmentExtractionError("PDF has too many pages")
        text_parts: list[str] = []
        warnings: list[str] = []
        for page in reader.pages:
            try:
                text_parts.append(page.extract_text() or "")
            except Exception:
                warnings.append("pdf_page_text_unavailable")
        text = "\n".join(text_parts).strip()
        if not text:
            warnings.append("empty_document")
        elif len(text) < self.MIN_TEXT_CHARACTERS:
            warnings.append("pdf_may_require_ocr")
        return AttachmentResult(
            filename=filename, mime_type=mime_type, document_type=DocumentType.PDF,
            text=text, warnings=warnings, metadata={"page_count": len(reader.pages)},
        )


class DocxExtractor:
    def extract(self, filename: str, mime_type: str, content: bytes) -> AttachmentResult:
        try:
            document = Document(BytesIO(content))
        except Exception as exc:
            raise AttachmentExtractionError("DOCX is damaged or unreadable") from exc
        text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text).strip()
        tables = [
            ExtractedTable(name=f"table_{index + 1}", rows=[[cell.text for cell in row.cells] for row in table.rows])
            for index, table in enumerate(document.tables)
        ]
        warnings = ["empty_document"] if not text and not tables else []
        return AttachmentResult(
            filename=filename, mime_type=mime_type, document_type=DocumentType.DOCX,
            text=text, tables=tables, warnings=warnings,
            metadata={"paragraph_count": len(document.paragraphs), "table_count": len(tables)},
        )


class XlsxExtractor:
    def __init__(self, *, max_sheets: int, max_rows: int, max_columns: int, max_uncompressed_bytes: int) -> None:
        self.max_sheets = max_sheets
        self.max_rows = max_rows
        self.max_columns = max_columns
        self.max_uncompressed_bytes = max_uncompressed_bytes

    def extract(self, filename: str, mime_type: str, content: bytes) -> AttachmentResult:
        self._validate_zip_limits(content)
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise AttachmentExtractionError("XLSX is damaged or unreadable") from exc
        try:
            if len(workbook.sheetnames) > self.max_sheets:
                raise AttachmentExtractionError("XLSX has too many worksheets")
            tables: list[ExtractedTable] = []
            warnings: list[str] = []
            text_parts: list[str] = []
            for worksheet in workbook.worksheets:
                rows: list[list[str | int | float | bool | None]] = []
                truncated = False
                for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                    if row_index > self.max_rows:
                        truncated = True
                        break
                    if len(row) > self.max_columns:
                        row = row[:self.max_columns]
                        truncated = True
                    normalized = [self._cell_value(value) for value in row]
                    rows.append(normalized)
                    text_parts.extend(str(value) for value in normalized if value is not None)
                if truncated:
                    warnings.append(f"xlsx_sheet_truncated:{worksheet.title}")
                tables.append(ExtractedTable(name=worksheet.title, rows=rows))
            text = "\n".join(text_parts).strip()
            if not text:
                warnings.append("empty_document")
            return AttachmentResult(
                filename=filename, mime_type=mime_type, document_type=DocumentType.XLSX,
                text=text, tables=tables, warnings=warnings,
                metadata={"sheet_count": len(workbook.sheetnames), "max_rows_per_sheet": self.max_rows, "max_columns": self.max_columns},
            )
        finally:
            workbook.close()

    def _validate_zip_limits(self, content: bytes) -> None:
        try:
            with ZipFile(BytesIO(content)) as archive:
                if len(archive.infolist()) > 1000:
                    raise AttachmentExtractionError("XLSX has too many archive entries")
                if sum(entry.file_size for entry in archive.infolist()) > self.max_uncompressed_bytes:
                    raise AttachmentExtractionError("XLSX expands beyond the allowed size")
        except BadZipFile as exc:
            raise AttachmentExtractionError("XLSX is damaged or unreadable") from exc

    @staticmethod
    def _cell_value(value: object) -> str | int | float | bool | None:
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        return str(value)


class ImageOcrExtractor:
    MAX_PIXELS = 25_000_000

    def extract(self, filename: str, mime_type: str, content: bytes) -> AttachmentResult:
        try:
            with Image.open(BytesIO(content)) as image:
                image.load()
                if image.width * image.height > self.MAX_PIXELS:
                    raise AttachmentExtractionError("Image has too many pixels")
                try:
                    text = pytesseract.image_to_string(image).strip()
                    warnings: list[str] = []
                except (pytesseract.TesseractNotFoundError, pytesseract.TesseractError, RuntimeError):
                    text = ""
                    warnings = ["ocr_failed"]
                if not text:
                    warnings.append("empty_document")
                return AttachmentResult(
                    filename=filename, mime_type=mime_type, document_type=DocumentType.IMAGE,
                    text=text, warnings=warnings,
                    metadata={"width": image.width, "height": image.height, "format": image.format},
                )
        except AttachmentExtractionError:
            raise
        except (UnidentifiedImageError, OSError, ValueError) as exc:
            raise AttachmentExtractionError("Image is damaged or unreadable") from exc
